"""
Lectura centralizada de inputs/Parametros_Sistema.xlsx.

POR QUE EXISTE ESTE MODULO
Hasta ahora, valores como los ajustes de la proteccion de cabecera (240 A,
dial 0.1, 1500 A...), la capacidad de corte de la subestacion (10 kA), los datos
de cortocircuito del equivalente de 115 kV o los nombres de las barras clave
estaban escritos directamente dentro de cada script. Eso tenia dos problemas:
el mismo numero aparecia repetido en varios archivos (con el riesgo de que uno
quedara desactualizado), y para saber de donde salia un valor habia que abrir el
codigo y buscarlo.

Ahora todo eso vive en un Excel de insumos, con su unidad, su descripcion y su
FUENTE documental, y los scripts lo leen de ahi. Para cambiar un ajuste se edita
el Excel, no el codigo.

Uso:
    from parametros import P
    P.cabecera("PICKUP_51_A")      -> 240.0
    P.proyecto("PROY_DIAL_51")     -> 0.05
    P.red_or("CAPACIDAD_CORTE_KA") -> 10.0
    P.elemento("BARRA_PC")         -> "P1 15344 13.2kV"
    P.anios()                      -> [(2026, 1.0), (2028, 1.0104)]

Si el Excel no existe o le falta una clave, se avisa con un mensaje que dice
exactamente que hoja y que fila hay que revisar, en vez de fallar con un
KeyError sin contexto.
"""

from pathlib import Path

import openpyxl

RUTA_PARAMETROS = Path(__file__).resolve().parent.parent / "inputs" / "Parametros_Sistema.xlsx"


class _Parametros:
    def __init__(self, ruta=RUTA_PARAMETROS):
        self.ruta = Path(ruta)
        self._cache = {}

    # ---- lectura basica ----
    def _hoja(self, nombre):
        if nombre in self._cache:
            return self._cache[nombre]
        if not self.ruta.exists():
            raise FileNotFoundError(
                f"No se encuentra el archivo de parametros '{self.ruta}'. "
                f"Es el Excel donde viven los ajustes de proteccion, los datos de la red del OR "
                f"y los nombres de los elementos del modelo."
            )
        wb = openpyxl.load_workbook(self.ruta, data_only=True)
        if nombre not in wb.sheetnames:
            raise KeyError(
                f"El archivo '{self.ruta.name}' no tiene la hoja '{nombre}'. "
                f"Hojas disponibles: {wb.sheetnames}"
            )
        ws = wb[nombre]
        datos = {}
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila and fila[0] is not None:
                datos[str(fila[0]).strip()] = fila[1]
        self._cache[nombre] = datos
        return datos

    def valor(self, hoja, clave):
        datos = self._hoja(hoja)
        if clave not in datos:
            raise KeyError(
                f"Falta el parametro '{clave}' en la hoja '{hoja}' de '{self.ruta.name}'. "
                f"Claves disponibles en esa hoja: {sorted(datos)}"
            )
        return datos[clave]

    # ---- accesos por hoja, para que el codigo que los usa se lea solo ----
    def cabecera(self, clave):
        """Ajustes de la proteccion de cabecera del circuito (datos del OR)."""
        return self.valor("Proteccion_Cabecera", clave)

    def proyecto(self, clave):
        """Ajustes propuestos para la proteccion del punto de conexion."""
        return self.valor("Proteccion_Proyecto", clave)

    def datos_proyecto(self, clave):
        """Datos de placa y caracteristicas del proyecto."""
        return self.valor("Datos_Proyecto", clave)

    def red_or(self, clave):
        """Datos de la red del Operador de Red y limites regulatorios."""
        return self.valor("Red_OR", clave)

    def elemento(self, clave):
        """Nombre exacto de un elemento del modelo de PowerFactory."""
        return self.valor("Elementos_Clave", clave)

    def despacho(self, clave):
        """Parametros del perfil de despacho solar."""
        return self.valor("Despacho_Solar", clave)

    def anios(self):
        """[(anio, factor_demanda), ...] del horizonte de analisis."""
        if not self.ruta.exists():
            raise FileNotFoundError(f"No se encuentra '{self.ruta}'.")
        wb = openpyxl.load_workbook(self.ruta, data_only=True)
        ws = wb["Anios_Analisis"]
        filas = []
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if fila and fila[0] is not None:
                filas.append((int(fila[0]), float(fila[1])))
        return filas

    def booleano(self, hoja, clave):
        """Lee un SI/NO del Excel como booleano."""
        valor = self.valor(hoja, clave)
        return str(valor).strip().upper() in ("SI", "SÍ", "TRUE", "1", "X")


P = _Parametros()
